"""
Automated tests for src/engine.py (the GUI-independent core of Schedule Optimizer).

Run from the repository root with the standard library only:

    python -m unittest discover -s tests -v

pytest also works if you have it installed:  pytest tests

Excel-export tests need openpyxl and are skipped automatically without it.
"""
import itertools
import random
import sys
import tempfile
import threading
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import engine  # noqa: E402

try:
    from openpyxl import load_workbook
except ImportError:  # export tests are skipped in that case
    load_workbook = None


# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
def make_rows(spec):
    """spec: (section, day, subject, part, start_text, end_text) tuples -> [Row]."""
    return [
        engine.Row(uid=i, section=sec, day=day, subject=subj, part=part,
                   start=engine.to_minutes(s), end=engine.to_minutes(e))
        for i, (sec, day, subj, part, s, e) in enumerate(spec)
    ]


def rows_from_sample(sample=None):
    sample = engine.SAMPLE if sample is None else sample
    rows, uid = [], 0
    for section, meetings in sample.items():
        for day, subject, part, start, end in meetings:
            rows.append(engine.Row(uid, section, day, subject, part,
                                   engine.to_minutes(start), engine.to_minutes(end)))
            uid += 1
    return rows


def run_search(rows, rankings=None, priority=("conflict", "days"), limit=10, **kwargs):
    keys, comps = engine.build_components(rows)
    rankings = {} if rankings is None else rankings
    best, nodes, cancelled = engine.search(comps, rankings, list(priority), limit, **kwargs)
    return keys, comps, best, nodes, cancelled


def idxs(best):
    return [entry[1] for entry in best]


def results_for(keys, comps, best, rankings):
    out = []
    for rank, entry in enumerate(best, 1):
        r = engine.describe(entry, keys, comps, rankings)
        r["rank"] = rank
        out.append(r)
    return out


# Two components (A, B), two sections each.  The four possible schedules are:
#   (0,0) A Mon 9-11 + B Mon 10-12 -> 60 min conflict, 1 day
#   (0,1) A Mon 9-11 + B Wed 9-11  ->  0 conflict, 2 days
#   (1,0) A Tue 9-11 + B Mon 10-12 ->  0 conflict, 2 days
#   (1,1) A Tue 9-11 + B Wed 9-11  ->  0 conflict, 2 days
def trade_off_rows(year="1"):
    a, b = f"{year}-A", f"{year}-B"
    return make_rows([
        (a, "Mon", "A", "LEC", "9:00AM", "11:00AM"),
        (b, "Tue", "A", "LEC", "9:00AM", "11:00AM"),
        (a, "Mon", "B", "LEC", "10:00AM", "12:00PM"),
        (b, "Wed", "B", "LEC", "9:00AM", "11:00AM"),
    ])


# ----------------------------------------------------------------------------
# Time parsing and formatting
# ----------------------------------------------------------------------------
class ToMinutesTests(unittest.TestCase):
    def test_accepted_formats(self):
        cases = {
            "7:00AM": 420, "7AM": 420, "7 am": 420, "7a": 420, "7P": 1140,
            "12:00AM": 0, "12AM": 0, "12:30AM": 30,
            "12:00PM": 720, "12PM": 720, "12:59PM": 779,
            "1:00PM": 780, "11:59PM": 1439, "2 p.m.": 840, "  9:30 am ": 570,
            "14:00": 840, "1430": 870, "730": 450, "0:00": 0, "00:00": 0, "23:59": 1439,
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(engine.to_minutes(text), expected)

    def test_bare_hour_is_24_hour_clock(self):
        self.assertEqual(engine.to_minutes("14"), 14 * 60)

    def test_rejected_formats(self):
        for text in ["", "abc", "AM", "24:00", "25:00", "13PM", "0AM", "0PM",
                     "7:60AM", "7:5AM", "7::00", "1:2:3", "7:00XM"]:
            with self.subTest(text=text):
                with self.assertRaises(ValueError):
                    engine.to_minutes(text)

    def test_error_message_names_the_bad_input(self):
        with self.assertRaisesRegex(ValueError, "Can't read the time 'abc'"):
            engine.to_minutes("abc")


class FormatTests(unittest.TestCase):
    def test_fmt(self):
        cases = {0: "12:00AM", 5: "12:05AM", 425: "7:05AM", 719: "11:59AM",
                 720: "12:00PM", 780: "1:00PM", 1439: "11:59PM"}
        for minutes, expected in cases.items():
            with self.subTest(minutes=minutes):
                self.assertEqual(engine.fmt(minutes), expected)

    def test_clock12_has_no_suffix(self):
        cases = {0: "12:00", 750: "12:30", 840: "2:00", 1439: "11:59"}
        for minutes, expected in cases.items():
            with self.subTest(minutes=minutes):
                self.assertEqual(engine.clock12(minutes), expected)

    def test_fmt_round_trips_through_to_minutes_for_every_minute_of_the_day(self):
        for minutes in range(24 * 60):
            self.assertEqual(engine.to_minutes(engine.fmt(minutes)), minutes)


# ----------------------------------------------------------------------------
# Small label / sorting helpers
# ----------------------------------------------------------------------------
class LabelHelperTests(unittest.TestCase):
    def test_year_of(self):
        cases = {"2-A": "2", "10-B": "10", "  3-D": "3", "12": "12",
                 "Z": "Other", "": "Other", "A2": "Other"}
        for section, expected in cases.items():
            with self.subTest(section=section):
                self.assertEqual(engine.year_of(section), expected)

    def test_year_label(self):
        self.assertEqual(engine.year_label("2"), "Year 2")
        self.assertEqual(engine.year_label("Other"), "Other sections")

    def test_year_sort_key_is_numeric_with_other_last(self):
        years = ["Other", "10", "2", "3"]
        self.assertEqual(sorted(years, key=engine.year_sort_key), ["2", "3", "10", "Other"])

    def test_natural_key_orders_numbers_numerically(self):
        sections = ["2-10", "2-2", "2-1", "1-B"]
        self.assertEqual(sorted(sections, key=engine.natural_key), ["1-B", "2-1", "2-2", "2-10"])

    def test_natural_key_ignores_case(self):
        self.assertEqual(engine.natural_key("b"), engine.natural_key("B"))

    def test_comp_label(self):
        self.assertEqual(engine.comp_label(("DSA", "ALL")), "DSA")
        self.assertEqual(engine.comp_label(("DSA", "LEC")), "DSA Lec")
        self.assertEqual(engine.comp_label(("DSA", "LAB")), "DSA Lab")

    def test_split_label(self):
        self.assertEqual(engine.split_label("DSA Lec [2-A]"), ("DSA Lec", "2-A"))
        self.assertEqual(engine.split_label("Foo [x] Lec [2-A]"), ("Foo [x] Lec", "2-A"))

    def test_split_label_inverts_label_construction(self):
        for key in [("DSA", "ALL"), ("DS and A", "LAB"), ("Intro", "LEC")]:
            label = f"{engine.comp_label(key)} [3-D]"
            self.assertEqual(engine.split_label(label), (engine.comp_label(key), "3-D"))

    def test_row_sort_key_orders_by_section_day_start_then_subject(self):
        rows = [
            engine.Row(0, "2-10", "Mon", "a", "LEC", 540, 600),
            engine.Row(1, "2-2", "Tue", "a", "LEC", 540, 600),
            engine.Row(2, "2-2", "Mon", "B", "LEC", 600, 660),
            engine.Row(3, "2-2", "Mon", "a", "LEC", 600, 660),    # subject order ignores case
            engine.Row(4, "2-2", "Mon", "z", "LEC", 480, 540),
        ]
        ordered = [r.uid for r in sorted(rows, key=engine.row_sort_key)]
        self.assertEqual(ordered, [4, 3, 2, 1, 0])


# ----------------------------------------------------------------------------
# Overlap arithmetic
# ----------------------------------------------------------------------------
class OverlapTests(unittest.TestCase):
    def test_partial_overlap(self):
        a = [("Mon", 540, 600, "a")]
        b = [("Mon", 570, 630, "b")]
        self.assertEqual(engine.overlap_minutes(a, b), 30)
        self.assertEqual(engine.overlap_minutes(b, a), 30)

    def test_touching_meetings_do_not_overlap(self):
        self.assertEqual(engine.overlap_minutes([("Mon", 540, 600, "a")], [("Mon", 600, 660, "b")]), 0)

    def test_different_days_do_not_overlap(self):
        self.assertEqual(engine.overlap_minutes([("Mon", 540, 600, "a")], [("Tue", 540, 600, "b")]), 0)

    def test_containment_counts_the_inner_length(self):
        self.assertEqual(engine.overlap_minutes([("Mon", 540, 720, "a")], [("Mon", 600, 660, "b")]), 60)

    def test_overlap_sums_over_every_meeting_pair(self):
        a = [("Mon", 540, 600, "a1"), ("Wed", 540, 600, "a2")]
        b = [("Mon", 570, 600, "b1"), ("Wed", 560, 580, "b2"), ("Fri", 540, 600, "b3")]
        self.assertEqual(engine.overlap_minutes(a, b), 30 + 20)

    def test_internal_overlap_sums_all_pairs(self):
        meetings = [("Mon", 540, 600, "a"), ("Mon", 570, 630, "b"), ("Mon", 590, 650, "c")]
        self.assertEqual(engine.internal_overlap(meetings), 30 + 10 + 40)

    def test_internal_overlap_of_zero_or_one_meeting_is_zero(self):
        self.assertEqual(engine.internal_overlap([]), 0)
        self.assertEqual(engine.internal_overlap([("Mon", 540, 600, "a")]), 0)

    def test_day_mask(self):
        self.assertEqual(engine.day_mask([]), 0)
        self.assertEqual(engine.day_mask([("Mon", 0, 1, "a"), ("Wed", 0, 1, "b"), ("Mon", 2, 3, "c")]), 0b000101)
        self.assertEqual(engine.day_mask([("Sat", 0, 1, "a")]), 0b100000)
        self.assertEqual(engine.popcount(0b101101), 4)


# ----------------------------------------------------------------------------
# Components and preference points
# ----------------------------------------------------------------------------
class BuildComponentsTests(unittest.TestCase):
    def setUp(self):
        self.rows = make_rows([
            ("1-B", "Mon", "Math", "LEC", "9:00AM", "10:00AM"),
            ("1-A", "Wed", "Math", "LEC", "1:00PM", "2:00PM"),
            ("1-A", "Tue", "Math", "LEC", "9:00AM", "10:00AM"),
            ("1-A", "Mon", "Math", "LAB", "1:00PM", "4:00PM"),
            ("1-A", "Thu", "Art", "ALL", "8:00AM", "9:00AM"),
        ])

    def test_keys_are_sorted_and_comps_line_up(self):
        keys, comps = engine.build_components(self.rows)
        self.assertEqual(keys, [("Art", "ALL"), ("Math", "LAB"), ("Math", "LEC")])
        self.assertEqual(len(comps), len(keys))

    def test_sections_grouped_in_natural_order_with_meetings_in_week_order(self):
        keys, comps = engine.build_components(self.rows)
        math_lec = comps[keys.index(("Math", "LEC"))]
        self.assertEqual([sec for sec, _ in math_lec], ["1-A", "1-B"])
        self.assertEqual(math_lec[0][1], [("Tue", 540, 600, "Math Lec [1-A]"),
                                          ("Wed", 780, 840, "Math Lec [1-A]")])
        self.assertEqual(math_lec[1][1], [("Mon", 540, 600, "Math Lec [1-B]")])

    def test_all_part_label_has_no_suffix(self):
        keys, comps = engine.build_components(self.rows)
        art = comps[keys.index(("Art", "ALL"))]
        self.assertEqual(art, [("1-A", [("Thu", 480, 540, "Art [1-A]")])])

    def test_excluded_components_are_dropped(self):
        keys, comps = engine.build_components(self.rows, excluded={("Math", "LAB"), ("Art", "ALL")})
        self.assertEqual(keys, [("Math", "LEC")])
        self.assertEqual(len(comps), 1)

    def test_no_rows_gives_nothing(self):
        self.assertEqual(engine.build_components([]), ([], []))

    def test_everything_excluded_gives_nothing(self):
        keys, comps = engine.build_components(self.rows, excluded={(r.subject, r.part) for r in self.rows})
        self.assertEqual((keys, comps), ([], []))


class MakePointsTests(unittest.TestCase):
    def test_points_follow_rank_position(self):
        points = engine.make_points({"3": ["3-D", "3-E"], "2": []})
        self.assertEqual(points("3-D"), 2)
        self.assertEqual(points("3-E"), 1)

    def test_unranked_sections_and_years_earn_nothing(self):
        points = engine.make_points({"3": ["3-D", "3-E"], "2": []})
        for section in ["3-F", "2-A", "4-A", "Z", ""]:
            with self.subTest(section=section):
                self.assertEqual(points(section), 0)


# ----------------------------------------------------------------------------
# Search: hand-built problems with known answers
# ----------------------------------------------------------------------------
class SearchOrderingTests(unittest.TestCase):
    def test_conflict_first_prefers_zero_conflict_then_breaks_ties_in_product_order(self):
        _, _, best, _, cancelled = run_search(trade_off_rows(), priority=["conflict", "days"])
        self.assertFalse(cancelled)
        self.assertEqual(idxs(best), [(0, 1), (1, 0), (1, 1), (0, 0)])

    def test_days_first_prefers_fewest_days_even_with_a_conflict(self):
        _, _, best, _, _ = run_search(trade_off_rows(), priority=["days", "conflict"])
        self.assertEqual(idxs(best), [(0, 0), (0, 1), (1, 0), (1, 1)])

    def test_limit_keeps_only_the_best_n(self):
        _, _, best, _, _ = run_search(trade_off_rows(), limit=1)
        self.assertEqual(idxs(best), [(0, 1)])
        _, _, best, _, _ = run_search(trade_off_rows(), limit=2)
        self.assertEqual(idxs(best), [(0, 1), (1, 0)])

    def test_limit_larger_than_search_space_returns_everything(self):
        _, _, best, _, _ = run_search(trade_off_rows(), limit=10_000)
        self.assertEqual(len(best), 4)

    def test_results_are_sorted_best_first_by_packed_key(self):
        _, _, best, _, _ = run_search(trade_off_rows(), limit=10)
        self.assertEqual([e[0] for e in best], sorted(e[0] for e in best))

    def test_single_component_single_section(self):
        rows = make_rows([("1-A", "Mon", "Solo", "LEC", "9:00AM", "10:00AM")])
        keys, _, best, nodes, cancelled = run_search(rows)
        self.assertEqual(keys, [("Solo", "LEC")])
        self.assertEqual(idxs(best), [(0,)])
        self.assertGreater(nodes, 0)
        self.assertFalse(cancelled)

    def test_six_day_schedule_is_not_confused_with_a_one_minute_conflict(self):
        # Option 0: one day, 1 minute of internal conflict.  Option 1: all six days, no conflict.
        # Whichever criterion is ranked first decides, even at the extremes of the day count.
        rows = make_rows(
            [("1-A", "Mon", "Solo", "LEC", "9:00AM", "10:00AM"),
             ("1-A", "Mon", "Solo", "LEC", "9:59AM", "10:59AM")]
            + [("1-B", day, "Solo", "LEC", "9:00AM", "10:00AM") for day in engine.DAY_ORDER])
        _, _, best, _, _ = run_search(rows, priority=["conflict", "days"])
        self.assertEqual(idxs(best), [(1,), (0,)])
        _, _, best, _, _ = run_search(rows, priority=["days", "conflict"])
        self.assertEqual(idxs(best), [(0,), (1,)])

    def test_overlap_inside_one_section_counts_as_conflict(self):
        # Section 1-A's own two meetings overlap by 30 min; 1-B has none.
        rows = make_rows([
            ("1-A", "Mon", "Solo", "LEC", "9:00AM", "10:00AM"),
            ("1-A", "Mon", "Solo", "LEC", "9:30AM", "10:30AM"),
            ("1-B", "Tue", "Solo", "LEC", "9:00AM", "10:00AM"),
        ])
        keys, comps, best, _, _ = run_search(rows, priority=["conflict", "days"])
        self.assertEqual(idxs(best), [(1,), (0,)])
        self.assertEqual(engine.describe(best[1], keys, comps, {})["conflict"], 0.5)


class SearchConstraintTests(unittest.TestCase):
    def test_max_conflict_minutes_excludes_worse_schedules(self):
        _, _, best, _, _ = run_search(trade_off_rows(), priority=["days", "conflict"], max_conflict_min=0)
        self.assertEqual(idxs(best), [(0, 1), (1, 0), (1, 1)])

    def test_max_conflict_boundary_is_inclusive(self):
        _, _, best, _, _ = run_search(trade_off_rows(), max_conflict_min=59)
        self.assertNotIn((0, 0), idxs(best))
        _, _, best, _, _ = run_search(trade_off_rows(), max_conflict_min=60)
        self.assertIn((0, 0), idxs(best))

    def test_max_days_excludes_schedules_spanning_more_days(self):
        _, _, best, _, _ = run_search(trade_off_rows(), max_days=1)
        self.assertEqual(idxs(best), [(0, 0)])

    def test_impossible_constraints_give_no_results(self):
        _, _, best, _, cancelled = run_search(trade_off_rows(), max_days=1, max_conflict_min=0)
        self.assertEqual(best, [])
        self.assertFalse(cancelled)


class SearchPreferenceTests(unittest.TestCase):
    RANK_A_FIRST = {"2": ["2-A", "2-B"]}

    def test_preference_first_accepts_a_conflict_to_get_the_favourite_section(self):
        keys, comps, best, _, _ = run_search(
            trade_off_rows("2"), rankings=self.RANK_A_FIRST, priority=["pref:2", "conflict", "days"])
        self.assertEqual(best[0][1], (0, 0))
        r = engine.describe(best[0], keys, comps, self.RANK_A_FIRST)
        self.assertEqual(r["prefs"], {"2": (4, 4)})
        self.assertEqual(r["conflict"], 1.0)

    def test_conflict_first_overrides_preference(self):
        keys, comps, best, _, _ = run_search(
            trade_off_rows("2"), rankings=self.RANK_A_FIRST, priority=["conflict", "pref:2", "days"])
        self.assertEqual(idxs(best), [(0, 1), (1, 0), (1, 1), (0, 0)])
        r = engine.describe(best[0], keys, comps, self.RANK_A_FIRST)
        self.assertEqual(r["prefs"], {"2": (3, 4)})
        self.assertEqual(r["conflict"], 0.0)

    def test_reversing_the_ranking_flips_the_choice(self):
        rankings = {"2": ["2-B", "2-A"]}
        _, _, best, _, _ = run_search(trade_off_rows("2"), rankings=rankings,
                                      priority=["conflict", "pref:2", "days"])
        self.assertEqual(best[0][1], (1, 1))

    def test_ranking_for_a_year_with_no_matching_sections_scores_zero(self):
        keys, comps, best, _, _ = run_search(trade_off_rows("1"), rankings={"2": ["2-A"]})
        self.assertEqual(engine.describe(best[0], keys, comps, {"2": ["2-A"]})["prefs"], {"2": (0, 0)})

    def test_empty_ranking_is_ignored(self):
        keys, comps, best, _, _ = run_search(trade_off_rows("1"), rankings={"1": []})
        self.assertEqual(engine.describe(best[0], keys, comps, {"1": []})["prefs"], {})


class SearchPriorityNormalisationTests(unittest.TestCase):
    def baseline(self):
        return idxs(run_search(trade_off_rows(), priority=["conflict", "days"])[2])

    def test_empty_priority_defaults_to_conflict_then_days(self):
        self.assertEqual(idxs(run_search(trade_off_rows(), priority=[])[2]), self.baseline())

    def test_missing_criteria_are_appended(self):
        a = idxs(run_search(trade_off_rows(), priority=["days"])[2])
        b = idxs(run_search(trade_off_rows(), priority=["days", "conflict"])[2])
        self.assertEqual(a, b)

    def test_unknown_inactive_and_duplicate_entries_are_ignored(self):
        priority = ["pref:9", "bogus", "conflict", "conflict", "days", "pref:1"]
        self.assertEqual(idxs(run_search(trade_off_rows(), priority=priority)[2]), self.baseline())

    def test_preference_for_an_empty_ranking_is_ignored(self):
        best = run_search(trade_off_rows(), rankings={"3": []}, priority=["pref:3", "conflict", "days"])[2]
        self.assertEqual(idxs(best), self.baseline())


# ----------------------------------------------------------------------------
# Search: progress reporting and cancellation
# ----------------------------------------------------------------------------
def wide_rows(subjects=8, sections=("1-A", "1-B", "1-C", "1-D")):
    """subjects x sections components -> a full tree of sum(4**k) nodes (87,381 for 8 x 4)."""
    spec = []
    for k in range(subjects):
        for s, sec in enumerate(sections):
            spec.append((sec, engine.DAY_ORDER[s], f"S{k}", "LEC", f"{8 + k}:00", f"{9 + k}:00"))
    return make_rows(spec)


class SearchStateAndCancelTests(unittest.TestCase):
    def test_state_reports_nodes_and_full_progress_when_finished(self):
        state = {"nodes": 0, "progress": 0.0}
        _, _, _, nodes, cancelled = run_search(trade_off_rows(), state=state)
        self.assertFalse(cancelled)
        self.assertEqual(state["nodes"], nodes)
        self.assertEqual(state["progress"], 1.0)

    def test_stop_event_cancels_search_and_keeps_partial_results(self):
        stop = threading.Event()
        stop.set()
        state = {"nodes": 0, "progress": 0.0}
        _, _, best, nodes, cancelled = run_search(wide_rows(), limit=10**6, stop=stop, state=state)
        self.assertTrue(cancelled)
        self.assertEqual(nodes, 0x4000)           # the flag is polled every 16,384 nodes
        self.assertEqual(state["nodes"], nodes)
        self.assertLess(state["progress"], 1.0)
        self.assertIsInstance(best, list)

    def test_unset_stop_event_does_not_cancel(self):
        stop = threading.Event()
        _, _, best, _, cancelled = run_search(wide_rows(), limit=5, stop=stop)
        self.assertFalse(cancelled)
        self.assertEqual(len(best), 5)


# ----------------------------------------------------------------------------
# Search: cross-check against an independent brute-force oracle
# ----------------------------------------------------------------------------
def oracle_order(priority, rankings):
    """Documented behaviour: valid listed criteria first, then any that are missing."""
    active = [y for y, r in rankings.items() if r]
    order = []
    for c in priority:
        valid = c in ("conflict", "days") or (c.startswith("pref:") and c[5:] in active)
        if valid and c not in order:
            order.append(c)
    for c in ("conflict", "days", *[f"pref:{y}" for y in active]):
        if c not in order:
            order.append(c)
    return order


def oracle_cost(comps, choice, rankings, order):
    meetings = [m for comp, a in zip(comps, choice) for m in comp[a][1]]
    conflict = 0
    for (d1, s1, e1, _), (d2, s2, e2, _) in itertools.combinations(meetings, 2):
        if d1 == d2:
            conflict += max(0, min(e1, e2) - max(s1, s2))
    days = len({m[0] for m in meetings})

    def pts(section, year):
        ranked = rankings.get(year) or []
        if engine.year_of(section) == year and section in ranked:
            return len(ranked) - ranked.index(section)
        return 0

    def deficit(year):
        total = 0
        for comp, a in zip(comps, choice):
            options = [pts(sec, year) for sec, _ in comp]
            total += max(options) - options[a]
        return total

    values = {"conflict": conflict, "days": days}
    for c in order:
        if c.startswith("pref:"):
            values[c] = deficit(c[5:])
    return tuple(values[c] for c in order), conflict, days


def oracle_best(comps, rankings, priority, limit, max_conflict_min=None, max_days=None):
    order = oracle_order(priority, rankings)
    found = []
    for choice in itertools.product(*[range(len(c)) for c in comps]):
        cost, conflict, days = oracle_cost(comps, choice, rankings, order)
        if max_conflict_min is not None and conflict > max_conflict_min:
            continue
        if max_days is not None and days > max_days:
            continue
        found.append((cost, choice))
    found.sort()
    return [choice for _, choice in found[:limit]]


def random_rows(rng):
    sections = ["1-A", "1-B", "2-A", "2-B"]
    spec = []
    for s in range(rng.randint(1, 4)):
        for part in rng.sample(engine.PARTS, rng.randint(1, 2)):
            for sec in rng.sample(sections, rng.randint(1, len(sections))):
                for _ in range(rng.randint(1, 2)):
                    start = rng.randrange(8 * 60, 18 * 60, 30)
                    end = start + rng.choice([60, 90, 120, 180])
                    spec.append((sec, rng.choice(engine.DAY_ORDER), f"Sub{s}", part,
                                 engine.fmt(start), engine.fmt(end)))
    return make_rows(spec)


class SearchMatchesBruteForceTests(unittest.TestCase):
    def test_random_problems(self):
        criteria = ["conflict", "days", "pref:1", "pref:2"]
        for seed in range(150):
            with self.subTest(seed=seed):
                rng = random.Random(seed)
                rows = random_rows(rng)
                rankings = {y: rng.sample([f"{y}-A", f"{y}-B"], rng.randint(0, 2)) for y in ("1", "2")}
                priority = rng.sample(criteria, rng.randint(0, len(criteria)))
                limit = rng.choice([1, 3, 10, 10_000])
                max_conflict = rng.choice([None, None, 0, 60, 120])
                max_days = rng.choice([None, None, 1, 2, 3])

                _, comps, best, _, cancelled = run_search(
                    rows, rankings, priority, limit,
                    max_conflict_min=max_conflict, max_days=max_days)
                expected = oracle_best(comps, rankings, priority, limit, max_conflict, max_days)

                self.assertFalse(cancelled)
                self.assertEqual(idxs(best), expected)

    def test_describe_agrees_with_oracle_on_conflict_days_and_preference(self):
        for seed in range(40):
            with self.subTest(seed=seed):
                rng = random.Random(1000 + seed)
                rows = random_rows(rng)
                rankings = {"1": rng.sample(["1-A", "1-B"], 2), "2": rng.sample(["2-A", "2-B"], 2)}
                priority = ["conflict", "days", "pref:1", "pref:2"]
                keys, comps, best, _, _ = run_search(rows, rankings, priority, limit=5)
                for entry in best:
                    r = engine.describe(entry, keys, comps, rankings)
                    _, conflict, days = oracle_cost(comps, entry[1], rankings, priority)
                    self.assertAlmostEqual(r["conflict"], conflict / 60)
                    self.assertEqual(r["days"], days)
                    self.assertEqual(len(r["assignment"]), len(keys))
                    for y, (got, top) in r["prefs"].items():
                        self.assertLessEqual(got, top)


# ----------------------------------------------------------------------------
# describe() and the text renderers
# ----------------------------------------------------------------------------
class DescribeTests(unittest.TestCase):
    def setUp(self):
        self.keys, self.comps, self.best, _, _ = run_search(trade_off_rows(), limit=10)

    def describe(self, idx, rankings=None):
        entry = next(e for e in self.best if e[1] == idx)
        return engine.describe(entry, self.keys, self.comps, rankings or {})

    def test_conflicting_schedule(self):
        r = self.describe((0, 0))
        self.assertEqual(r["conflict"], 1.0)
        self.assertEqual(r["days"], 1)
        self.assertEqual(r["days_used"], ["Mon"])
        self.assertEqual(r["prefs"], {})
        self.assertEqual(r["assignment"], [(("A", "LEC"), "1-A"), (("B", "LEC"), "1-A")])
        self.assertEqual(len(r["meetings"]), 2)

    def test_days_used_follow_the_week_not_the_alphabet(self):
        rows = make_rows([("1-A", "Fri", "Solo", "LEC", "9:00AM", "10:00AM"),
                          ("1-A", "Mon", "Solo", "LEC", "9:00AM", "10:00AM"),
                          ("1-A", "Thu", "Solo", "LEC", "9:00AM", "10:00AM")])
        keys, comps, best, _, _ = run_search(rows)
        r = engine.describe(best[0], keys, comps, {})
        self.assertEqual(r["days_used"], ["Mon", "Thu", "Fri"])
        self.assertEqual(r["days"], 3)

    def test_clean_schedule_lists_days_in_week_order(self):
        r = self.describe((1, 1))
        self.assertEqual(r["conflict"], 0.0)
        self.assertEqual(r["days_used"], ["Tue", "Wed"])
        self.assertEqual(r["assignment"], [(("A", "LEC"), "1-B"), (("B", "LEC"), "1-B")])


class SummaryAndTimetableTests(unittest.TestCase):
    def test_summary_line_without_preferences(self):
        r = {"conflict": 0, "days": 1, "days_used": ["Mon"], "prefs": {}}
        self.assertEqual(engine.summary_line(r), "Conflict: 0.0 hr | Days: 1 (Mon)")

    def test_summary_line_sorts_years_numerically_with_other_last(self):
        r = {"conflict": 1.5, "days": 2, "days_used": ["Mon", "Tue"],
             "prefs": {"Other": (0, 0), "10": (3, 4), "2": (1, 2)}}
        self.assertEqual(
            engine.summary_line(r),
            "Conflict: 1.5 hr | Days: 2 (Mon, Tue) | Year 2 preference: 1/2"
            " | Year 10 preference: 3/4 | Other sections preference: 0/0")

    def test_choice_note(self):
        rankings = {"3": ["3-D", "3-E"]}
        self.assertEqual(engine.choice_note("3-D", rankings), "   (your choice #1)")
        self.assertEqual(engine.choice_note("3-E", rankings), "   (your choice #2)")
        self.assertEqual(engine.choice_note("3-X", rankings), "")
        self.assertEqual(engine.choice_note("2-A", rankings), "")

    def test_timetable_lines_flag_only_overlapping_meetings(self):
        meetings = [
            ("Wed", 540, 600, "D [1-A]"),
            ("Mon", 540, 600, "A [1-A]"),
            ("Mon", 570, 630, "B [1-A]"),
            ("Mon", 630, 690, "C [1-A]"),   # starts exactly when B ends: no conflict
        ]
        lines = engine.timetable_lines(meetings)

        def line_for(label):
            return next(line for line in lines if label in line)

        self.assertEqual(lines[0], "  Mon:")
        self.assertIn("9:00AM - 10:00AM", line_for("A [1-A]"))
        self.assertTrue(line_for("A [1-A]").endswith("<-- CONFLICT"))
        self.assertTrue(line_for("B [1-A]").endswith("<-- CONFLICT"))
        self.assertNotIn("CONFLICT", line_for("C [1-A]"))
        self.assertNotIn("CONFLICT", line_for("D [1-A]"))
        self.assertNotIn("  Tue:", lines)
        self.assertLess(lines.index("  Mon:"), lines.index("  Wed:"))

    def test_timetable_lines_flag_a_meeting_fully_inside_another(self):
        lines = engine.timetable_lines([("Fri", 540, 720, "Big [1-A]"), ("Fri", 600, 660, "Small [1-A]")])
        self.assertTrue(all(l.endswith("<-- CONFLICT") for l in lines if "[1-A]" in l))


# ----------------------------------------------------------------------------
# Excel export
# ----------------------------------------------------------------------------
@unittest.skipUnless(engine.HAVE_XLSX and load_workbook, "openpyxl is not installed")
class ExportXlsxTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.path = str(Path(self._tmp.name) / "out.xlsx")

    def export(self, results):
        engine.export_xlsx(results, self.path)
        return load_workbook(self.path)

    @staticmethod
    def custom_result(meetings, rank=1):
        return {"rank": rank, "meetings": meetings, "conflict": 0.0, "days": 1, "days_used": ["Mon"],
                "prefs": {}, "assignment": [(("X", "LEC"), "1-A")]}

    def trade_off_results(self, priority=("conflict", "days"), limit=10):
        keys, comps, best, _, _ = run_search(trade_off_rows(), priority=priority, limit=limit)
        return results_for(keys, comps, best, {})

    def test_one_sheet_per_result_named_by_rank(self):
        wb = self.export(self.trade_off_results())
        self.assertEqual(wb.sheetnames, ["#1", "#2", "#3", "#4"])

    def test_header_row_has_weekday_names_and_no_saturday_by_default(self):
        ws = self.export(self.trade_off_results())["#1"]
        self.assertEqual(ws["A1"].value, "Class Schedule")
        self.assertEqual([ws.cell(1, c).value for c in range(3, 8)],
                         ["MONDAY", "TUESDAY", "WEDNESDAY", "THURSDAY", "FRIDAY"])
        self.assertIsNone(ws["H1"].value)

    def test_saturday_column_appears_only_when_needed(self):
        result = self.custom_result([("Sat", 540, 600, "X Lec [1-A]")])
        ws = self.export([result])["#1"]
        self.assertEqual(ws["H1"].value, "SATURDAY")

    def test_time_column_covers_7am_to_9pm_in_half_hour_rows(self):
        ws = self.export(self.trade_off_results())["#1"]
        self.assertEqual((ws["A2"].value, ws["B2"].value), ("7:00", "7:30"))
        self.assertEqual((ws["A29"].value, ws["B29"].value), ("8:30", "9:00"))

    def test_class_block_is_merged_and_labelled(self):
        ws = self.export(self.trade_off_results())["#1"]    # rank 1 = A 1-A (Mon), B 1-B (Wed)
        merged = {str(r) for r in ws.merged_cells.ranges}
        self.assertIn("C6:C9", merged)                         # Mon 9:00-11:00
        self.assertEqual(ws["C6"].value, "A Lec\nSection 1-A")
        self.assertEqual(ws["E6"].value, "B Lec\nSection 1-B")  # Wed 9:00-11:00
        self.assertTrue(ws["C6"].fill.start_color.rgb.endswith("A9D99B"))
        self.assertTrue(ws["D6"].fill.start_color.rgb.endswith("FFEFBF"))   # Tuesday is empty

    def test_conflicting_meetings_share_one_red_cell_marked_conflict(self):
        results = self.trade_off_results(priority=["days", "conflict"])    # rank 1 = the 60-minute conflict
        ws = self.export(results)["#1"]
        text = ws["C6"].value
        self.assertTrue(text.startswith("CONFLICT"))
        self.assertIn("A Lec - 1-A (9:00-11:00)", text)
        self.assertIn("B Lec - 1-A (10:00-12:00)", text)
        self.assertTrue(ws["C6"].fill.start_color.rgb.endswith("F4A6A6"))

    def test_back_to_back_slot_aligned_classes_stay_in_separate_blocks(self):
        result = self.custom_result([("Mon", 540, 600, "X Lec [1-A]"), ("Mon", 600, 660, "Y Lec [1-A]")])
        ws = self.export([result])["#1"]
        merged = {str(r) for r in ws.merged_cells.ranges}
        self.assertTrue({"C6:C7", "C8:C9"} <= merged)
        self.assertEqual(ws["C6"].value, "X Lec\nSection 1-A")
        self.assertEqual(ws["C8"].value, "Y Lec\nSection 1-A")

    def test_back_to_back_meetings_sharing_a_slot_are_not_marked_as_conflict(self):
        result = self.custom_result([("Mon", 540, 585, "X Lec [1-A]"), ("Mon", 585, 630, "Y Lec [1-A]")])
        ws = self.export([result])["#1"]
        self.assertFalse(ws["C6"].value.startswith("CONFLICT"))
        self.assertIn("X Lec - 1-A (9:00-9:45)", ws["C6"].value)
        self.assertIn("Y Lec - 1-A (9:45-10:30)", ws["C6"].value)
        self.assertTrue(ws["C6"].fill.start_color.rgb.endswith("A9D99B"))

    def test_summary_footer_sits_below_the_grid(self):
        ws = self.export(self.trade_off_results())["#1"]
        self.assertEqual(ws["A31"].value, "Rank #1")
        self.assertTrue(ws["A32"].value.startswith("Conflict: 0.0 hr | Days: 2"))
        self.assertEqual(ws["A33"].value, "Sections: A Lec -> 1-A, B Lec -> 1-B")

    def test_grid_grows_for_early_and_late_classes(self):
        result = self.custom_result([("Mon", 360, 420, "Early [1-A]"), ("Mon", 1290, 1335, "Late [1-A]")])
        ws = self.export([result])["#1"]
        self.assertEqual(ws["A2"].value, "6:00")                # grid starts at 6:00 AM
        # 6:00 -> 22:30 is 33 half-hour slots, so the footer moves from row 31 to row 36
        self.assertEqual(ws["A36"].value, "Rank #1")

    def test_missing_openpyxl_gives_a_helpful_error(self):
        with mock.patch.object(engine, "HAVE_XLSX", False):
            with self.assertRaisesRegex(RuntimeError, "openpyxl"):
                engine.export_xlsx([], self.path)


# ----------------------------------------------------------------------------
# The built-in sample data (end to end)
# ----------------------------------------------------------------------------
class SampleDataTests(unittest.TestCase):
    def test_sample_is_well_formed(self):
        for section, meetings in engine.SAMPLE.items():
            for day, subject, part, start, end in meetings:
                with self.subTest(section=section, subject=subject, day=day):
                    self.assertIn(day, engine.DAY_ORDER)
                    self.assertIn(part, engine.PARTS)
                    self.assertLess(engine.to_minutes(start), engine.to_minutes(end))

    def test_sample_rankings_and_priority_refer_to_real_things(self):
        for year, ranked in engine.SAMPLE_RANKINGS.items():
            for section in ranked:
                self.assertIn(section, engine.SAMPLE)
                self.assertEqual(engine.year_of(section), year)
        for criterion in engine.SAMPLE_PRIORITY:
            self.assertTrue(criterion in ("conflict", "days") or criterion[5:] in engine.SAMPLE_RANKINGS)

    def test_end_to_end_search_on_the_sample(self):
        rows = rows_from_sample()
        keys, comps = engine.build_components(rows)
        best, nodes, cancelled = engine.search(
            comps, engine.SAMPLE_RANKINGS, engine.SAMPLE_PRIORITY, 5)
        self.assertFalse(cancelled)
        self.assertGreater(nodes, 0)
        self.assertEqual(len(best), 5)

        results = results_for(keys, comps, best, engine.SAMPLE_RANKINGS)
        for r in results:
            self.assertEqual([k for k, _ in r["assignment"]], keys)    # every component is scheduled
            self.assertEqual(r["days"], len(r["days_used"]))

        # Results come out best-first under the sample priority (conflict, days, year-3 pref, year-2 pref).
        def cost(r):
            return (r["conflict"], r["days"],
                    r["prefs"]["3"][1] - r["prefs"]["3"][0],
                    r["prefs"]["2"][1] - r["prefs"]["2"][0])
        costs = [cost(r) for r in results]
        self.assertEqual(costs, sorted(costs))


if __name__ == "__main__":
    unittest.main()
