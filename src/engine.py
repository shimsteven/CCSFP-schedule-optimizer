"""
Schedule Optimizer - engine
===========================

GUI-independent core shared by both front-ends (schedule_optimizer_ctk.py and
schedule_optimizer_gui.py): time parsing, the branch-and-bound search, result
description and Excel export.  Nothing in here imports a GUI toolkit.
"""
from __future__ import annotations

import bisect
import itertools
import re
from collections import defaultdict
from dataclasses import dataclass

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    HAVE_XLSX = True
except ImportError:  # the app still works; only Excel export is disabled
    HAVE_XLSX = False

# ----------------------------------------------------------------------------
# Basics
# ----------------------------------------------------------------------------
DAY_ORDER = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]
DAY_BIT = {d: 1 << i for i, d in enumerate(DAY_ORDER)}
PARTS = ["LEC", "LAB", "ALL"]
popcount = getattr(int, "bit_count", lambda x: bin(x).count("1"))

SAMPLE = {
    "2-A": [("Mon", "Mbapp", "LAB", "2:00PM", "5:00PM"), ("Mon", "DSA", "LEC", "5:30PM", "7:30PM"),
            ("Thu", "InfoMan", "LAB", "7:00AM", "10:00AM"), ("Thu", "Mbapp", "LEC", "10:30AM", "12:30PM"),
            ("Fri", "InfoMan", "LEC", "11:30AM", "1:30PM"), ("Fri", "DSA", "LAB", "2:00PM", "5:00PM")],
    "2-B": [("Mon", "InfoMan", "LEC", "11:30AM", "1:30PM"), ("Mon", "DSA", "LAB", "2:00PM", "5:00PM"),
            ("Tue", "Mbapp", "LAB", "2:00PM", "5:00PM"), ("Tue", "DSA", "LEC", "5:30PM", "7:30PM"),
            ("Fri", "InfoMan", "LAB", "7:00AM", "10:00AM"), ("Fri", "Mbapp", "LEC", "10:30AM", "12:30PM")],
    "2-C": [("Mon", "InfoMan", "LAB", "7:00AM", "10:00AM"), ("Mon", "Mbapp", "LEC", "10:30AM", "12:30PM"),
            ("Tue", "InfoMan", "LEC", "11:30AM", "1:30PM"), ("Tue", "DSA", "LAB", "2:00PM", "5:00PM"),
            ("Wed", "Mbapp", "LAB", "2:00PM", "5:00PM"), ("Wed", "DSA", "LEC", "5:00PM", "7:30PM")],
    "2-D": [("Tue", "InfoMan", "LAB", "7:00AM", "10:00AM"), ("Tue", "Mbapp", "LEC", "10:30AM", "12:30PM"),
            ("Wed", "InfoMan", "LEC", "11:30AM", "1:30PM"), ("Wed", "DSA", "LAB", "2:00PM", "5:00PM"),
            ("Thu", "Mbapp", "LAB", "2:00PM", "5:00PM"), ("Thu", "DSA", "LEC", "5:00PM", "7:00PM")],
    "3-D": [("Mon", "DS and A", "LEC", "3:30PM", "5:30PM"), ("Mon", "AdvNet", "LAB", "6:00PM", "9:00PM"),
            ("Tue", "AdvNet", "LEC", "3:30PM", "5:30PM"), ("Wed", "ArtApp", "ALL", "3:00PM", "4:30PM"),
            ("Wed", "QA", "ALL", "4:30PM", "6:00PM"), ("Thu", "QA", "ALL", "2:00PM", "3:30PM"),
            ("Thu", "BPM", "ALL", "6:00PM", "7:30PM"), ("Fri", "DS and A", "LAB", "11:30AM", "2:30PM"),
            ("Fri", "ArtApp", "ALL", "3:30PM", "5:00PM"), ("Fri", "BPM", "ALL", "6:00PM", "7:30PM")],
    "3-E": [("Mon", "ArtApp", "ALL", "2:30PM", "5:30PM"), ("Mon", "BPM", "ALL", "6:00PM", "7:30PM"),
            ("Tue", "BPM", "ALL", "6:00PM", "7:30PM"), ("Tue", "QA", "ALL", "7:30PM", "9:00PM"),
            ("Thu", "DS and A", "LEC", "3:30PM", "5:30PM"), ("Thu", "AdvNet", "LAB", "6:00PM", "9:00PM"),
            ("Fri", "AdvNet", "LEC", "3:30PM", "5:30PM")],
}
SAMPLE_RANKINGS = {"3": ["3-D"], "2": ["2-B"]}
SAMPLE_PRIORITY = ["conflict", "days", "pref:3", "pref:2"]

_TIME_RE = re.compile(r"^(\d{1,2})(?::?(\d{2}))?(AM|PM|A|P)?$")


def to_minutes(text):
    """'7:00AM', '2 pm', '2PM', '14:00', '1430' -> minutes after midnight."""
    t = str(text).strip().upper().replace(" ", "").replace(".", "")
    m = _TIME_RE.match(t)
    err = ValueError(f"Can't read the time '{text}'. Use something like 7:00AM, 2 PM or 14:00.")
    if not m:
        raise err
    hour, minute, suffix = int(m.group(1)), int(m.group(2) or 0), m.group(3)
    if minute > 59:
        raise err
    if suffix:
        if not 1 <= hour <= 12:
            raise err
        hour %= 12
        if suffix.startswith("P"):
            hour += 12
    elif hour > 23:
        raise err
    return hour * 60 + minute


def fmt(minutes):
    h, m = divmod(minutes, 60)
    return f"{h % 12 or 12}:{m:02d}{'AM' if h < 12 else 'PM'}"


def clock12(minutes):
    h, m = divmod(minutes, 60)
    return f"{h % 12 or 12}:{m:02d}"


def year_of(section):
    m = re.match(r"\s*(\d+)", section)
    return m.group(1) if m else "Other"


def year_label(y):
    return f"Year {y}" if y.isdigit() else "Other sections"


def year_sort_key(y):
    return (0, int(y)) if y.isdigit() else (1, 0)


def natural_key(s):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", s)]


def comp_label(key):
    subject, part = key
    return subject if part == "ALL" else f"{subject} {part.title()}"


@dataclass
class Row:
    """One weekly class meeting of one section."""
    uid: int
    section: str
    day: str
    subject: str
    part: str
    start: int
    end: int


def row_sort_key(r):
    return (natural_key(r.section), DAY_ORDER.index(r.day), r.start, r.subject.lower())


# ----------------------------------------------------------------------------
# Engine
# ----------------------------------------------------------------------------
class Cancelled(Exception):
    pass


def build_components(rows, excluded=()):
    """-> (keys, comps).  comps[k] = [(section, [(day, start, end, label)]), ...]"""
    options = defaultdict(lambda: defaultdict(list))
    for r in sorted(rows, key=row_sort_key):
        key = (r.subject, r.part)
        if key in excluded:
            continue
        options[key][r.section].append((r.day, r.start, r.end, f"{comp_label(key)} [{r.section}]"))
    keys = sorted(options)
    return keys, [list(options[k].items()) for k in keys]


def overlap_minutes(m1, m2):
    total = 0
    for d1, s1, e1, _ in m1:
        for d2, s2, e2, _ in m2:
            if d1 == d2:
                o = min(e1, e2) - max(s1, s2)
                if o > 0:
                    total += o
    return total


def internal_overlap(meetings):
    return sum(overlap_minutes([a], [b]) for a, b in itertools.combinations(meetings, 2))


def day_mask(meetings):
    mask = 0
    for m in meetings:
        mask |= DAY_BIT[m[0]]
    return mask


def make_points(rankings):
    """Points a section earns from the user's ranking of its year level."""
    def points(section):
        ranked = rankings.get(year_of(section))
        if ranked and section in ranked:
            return len(ranked) - ranked.index(section)
        return 0
    return points


def search(comps, rankings, priority, limit, max_conflict_min=None, max_days=None,
           stop=None, state=None):
    """
    Branch-and-bound over every way of assigning each component to one section.
    Returns (best, nodes, cancelled) where best = sorted [(key, idxs)], best first.

    Every ranking criterion is turned into a 'cost' that can only grow as more
    components are assigned (conflict minutes, days used, preference points lost).
    The costs are packed into ONE integer in priority order (mixed radix), so
    comparing two schedules is a single integer comparison and the running cost
    of a partial schedule is a valid lower bound for pruning.
    """
    n = len(comps)
    points = make_points(rankings)
    active_years = [y for y, r in rankings.items() if r]

    order = []
    for c in priority:
        ok = c in ("conflict", "days") or (c.startswith("pref:") and c[5:] in active_years)
        if ok and c not in order:
            order.append(c)
    for c in ("conflict", "days", *[f"pref:{y}" for y in active_years]):
        if c not in order:
            order.append(c)
    pyears = [c[5:] for c in order if c.startswith("pref:")]

    base = [[internal_overlap(m) for _, m in c] for c in comps]
    masks = [[day_mask(m) for _, m in c] for c in comps]
    pair = [
        [
            [[overlap_minutes(comps[i][ci][1], comps[k][a][1]) for a in range(len(comps[k]))]
             for ci in range(len(comps[i]))]
            for i in range(k)
        ]
        for k in range(n)
    ]
    # preference "deficit" = points lost versus the best option of that component
    pdef = {}
    for y in pyears:
        rows_ = []
        for k in range(n):
            pts = [points(sec) if year_of(sec) == y else 0 for sec, _ in comps[k]]
            mx = max(pts)
            rows_.append([mx - p for p in pts])
        pdef[y] = rows_

    conflict_bound = 1 + sum(
        max(base[k]) + sum(max(max(row) for row in pair[k][i]) for i in range(k))
        for k in range(n)
    )
    radix = {"conflict": conflict_bound, "days": len(DAY_ORDER) + 1}
    for y in pyears:
        radix[f"pref:{y}"] = 1 + sum(max(row) for row in pdef[y])
    weight, w = {}, 1
    for c in reversed(order):
        weight[c] = w
        w *= radix[c]
    wc, wd = weight["conflict"], weight["days"]
    prefterm = [
        [sum(pdef[y][k][a] * weight[f"pref:{y}"] for y in pyears) for a in range(len(comps[k]))]
        for k in range(n)
    ]

    max_c = float("inf") if max_conflict_min is None else max_conflict_min
    max_d = len(DAY_ORDER) if max_days is None else max_days
    len0 = len(comps[0])
    len1 = len(comps[1]) if n > 1 else 1

    best, idx, nodes = [], [0] * n, 0

    def dfs(k, acc, conflict, mask):
        nonlocal nodes
        nodes += 1
        if not nodes & 0x3FFF:
            if state is not None:
                state["nodes"] = nodes
            if stop is not None and stop.is_set():
                raise Cancelled
        key = acc + popcount(mask) * wd
        if len(best) == limit and key >= best[-1][0]:   # ties lose (later in product order)
            return
        if k == n:
            bisect.insort(best, (key, tuple(idx)))
            if len(best) > limit:
                best.pop()
            return
        prev_rows, ck, bk, mk, pk = pair[k], comps[k], base[k], masks[k], prefterm[k]
        for a in range(len(ck)):
            c = conflict + bk[a]
            for i in range(k):
                c += prev_rows[i][idx[i]][a]
            if c > max_c:
                continue
            m = mask | mk[a]
            if popcount(m) > max_d:
                continue
            idx[k] = a
            if k < 2 and state is not None:
                state["progress"] = (idx[0] + (a / len1 if k == 1 else 0)) / len0
            dfs(k + 1, acc + (c - conflict) * wc + pk[a], c, m)

    cancelled = False
    try:
        dfs(0, 0, 0, 0)
    except Cancelled:
        cancelled = True
    if state is not None:
        state["nodes"] = nodes
        if not cancelled:
            state["progress"] = 1.0
    return best, nodes, cancelled


def describe(entry, keys, comps, rankings):
    points = make_points(rankings)
    assignment, meetings = [], []
    for k, a in enumerate(entry[1]):
        sec, mtgs = comps[k][a]
        assignment.append((keys[k], sec))
        meetings.extend(mtgs)
    conflict = internal_overlap(meetings) / 60
    days_used = sorted({m[0] for m in meetings}, key=DAY_ORDER.index)
    prefs = {}
    for y, ranked in rankings.items():
        if not ranked:
            continue
        got = sum(points(sec) for _, sec in assignment if year_of(sec) == y)
        top = sum(max(points(s) if year_of(s) == y else 0 for s, _ in comps[k])
                  for k in range(len(comps)))
        prefs[y] = (got, top)
    return {"conflict": conflict, "days": len(days_used), "days_used": days_used, "prefs": prefs,
            "assignment": assignment, "meetings": meetings}


def split_label(label):
    name, _, sec = label.rpartition(" [")
    return name, sec.rstrip("]")


def timetable_lines(meetings):
    by_day = defaultdict(list)
    for day, start, end, label in meetings:
        by_day[day].append((start, end, label))
    out = []
    for day in DAY_ORDER:
        if day not in by_day:
            continue
        out.append(f"  {day}:")
        slots = sorted(by_day[day])
        for i, (start, end, label) in enumerate(slots):
            clash = any(j != i and min(end, e2) > max(start, s2)
                        for j, (s2, e2, _) in enumerate(slots))
            out.append(f"    {fmt(start):>8} - {fmt(end):<8} {label}{'   <-- CONFLICT' if clash else ''}")
    return out


def choice_note(section, rankings):
    ranked = rankings.get(year_of(section))
    return f"   (your choice #{ranked.index(section) + 1})" if ranked and section in ranked else ""


def summary_line(r):
    s = f"Conflict: {r['conflict']:.1f} hr | Days: {r['days']} ({', '.join(r['days_used'])})"
    for y in sorted(r["prefs"], key=year_sort_key):
        got, top = r["prefs"][y]
        s += f" | {year_label(y)} preference: {got}/{top}"
    return s


# ----------------------------------------------------------------------------
# Excel export (one sheet per schedule, weekly grid)
# ----------------------------------------------------------------------------
def export_xlsx(results, path):
    if not HAVE_XLSX:
        raise RuntimeError("Excel export needs openpyxl.  Install it with:  pip install openpyxl")
    grid_start, grid_end, slot = 7 * 60, 21 * 60, 30
    full_day = {"Mon": "MONDAY", "Tue": "TUESDAY", "Wed": "WEDNESDAY",
                "Thu": "THURSDAY", "Fri": "FRIDAY", "Sat": "SATURDAY"}
    font = "Arial"
    fill_header = PatternFill("solid", start_color="FFFF66")
    fill_time = PatternFill("solid", start_color="A8D5E8")
    fill_empty = PatternFill("solid", start_color="FFEFBF")
    fill_class = PatternFill("solid", start_color="A9D99B")
    fill_conflict = PatternFill("solid", start_color="F4A6A6")
    side = Side(style="thin", color="BFBFBF")
    border = Border(left=side, right=side, top=side, bottom=side)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def write_sheet(ws, r):
        meetings = r["meetings"]
        days = DAY_ORDER[:5] + (["Sat"] if any(m[0] == "Sat" for m in meetings) else [])
        lo = min(grid_start, min(m[1] for m in meetings) // slot * slot)
        hi = max(grid_end, -(-max(m[2] for m in meetings) // slot) * slot)
        n_slots = (hi - lo) // slot
        last_col = 2 + len(days)

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
        ws.cell(1, 1, "Class Schedule")
        for c in range(1, last_col + 1):
            cell = ws.cell(1, c)
            cell.fill, cell.border, cell.alignment = fill_header, border, center
            cell.font = Font(name=font, size=11, bold=(c <= 2))
        for j, day in enumerate(days):
            ws.cell(1, 3 + j, full_day[day])

        for s in range(n_slots):
            row, t0 = 2 + s, lo + s * slot
            for c, text in ((1, clock12(t0)), (2, clock12(t0 + slot))):
                cell = ws.cell(row, c, text)
                cell.fill, cell.border, cell.alignment = fill_time, border, center
                cell.font = Font(name=font, size=11)
            for j in range(len(days)):
                cell = ws.cell(row, 3 + j)
                cell.fill, cell.border = fill_empty, border
            ws.row_dimensions[row].height = 21

        for j, day in enumerate(days):
            col = 3 + j
            todays = sorted((m for m in meetings if m[0] == day), key=lambda m: (m[1], m[2]))
            clusters = []
            for m in todays:
                rs, re_ = m[1] // slot * slot, -(-m[2] // slot) * slot
                if clusters and rs < clusters[-1]["end"]:
                    clusters[-1]["items"].append(m)
                    clusters[-1]["end"] = max(clusters[-1]["end"], re_)
                else:
                    clusters.append({"start": rs, "end": re_, "items": [m]})
            for cl in clusters:
                r0 = 2 + (cl["start"] - lo) // slot
                r1 = 2 + (cl["end"] - lo) // slot - 1
                items = cl["items"]
                clash = any(overlap_minutes([a], [b]) > 0 for a, b in itertools.combinations(items, 2))
                if len(items) == 1:
                    name, sec = split_label(items[0][3])
                    text = f"{name}\nSection {sec}"
                else:
                    parts = []
                    for _, s_, e_, lab in items:
                        name, sec = split_label(lab)
                        parts.append(f"{name} - {sec} ({clock12(s_)}-{clock12(e_)})")
                    text = ("CONFLICT\n" if clash else "") + "\n".join(parts)
                if r1 > r0:
                    ws.merge_cells(start_row=r0, start_column=col, end_row=r1, end_column=col)
                for rr in range(r0, r1 + 1):
                    ws.cell(rr, col).fill = fill_conflict if clash else fill_class
                    ws.cell(rr, col).border = border
                top = ws.cell(r0, col, text)
                top.alignment = center
                top.font = Font(name=font, size=11, bold=clash)

        ws.column_dimensions["A"].width = 8.5
        ws.column_dimensions["B"].width = 8.5
        for j in range(len(days)):
            ws.column_dimensions[get_column_letter(3 + j)].width = 24
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        row = 3 + n_slots
        lines = [f"Rank #{r['rank']}", summary_line(r),
                 "Sections: " + ", ".join(f"{comp_label(k)} -> {sec}" for k, sec in r["assignment"])]
        for i, line in enumerate(lines):
            cell = ws.cell(row + i, 1, line)
            cell.font = Font(name=font, size=10, bold=(i == 0))
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            ws.merge_cells(start_row=row + i, start_column=1, end_row=row + i, end_column=last_col)
            ws.row_dimensions[row + i].height = 15 if i < 2 else 45

    wb = Workbook()
    wb.remove(wb.active)
    for r in results:
        write_sheet(wb.create_sheet(f"#{r['rank']}"), r)
    wb.save(path)
